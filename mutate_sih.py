from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# MUST use the file specified by the user
prs = Presentation('SIH2026_Final_Presentation.pptx')

# We assume the slide order is standard SIH 6-slide template.
# If Slide 7 exists, we will delete it later.

def format_text(shape, text, font_size=None, bold=None, align=None, color_rgb=None):
    if shape.has_text_frame:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = text
        if font_size: p.font.size = Pt(font_size)
        if bold is not None: p.font.bold = bold
        if align: p.alignment = align
        if color_rgb: p.font.color.rgb = RGBColor(*color_rgb)

def clean_footer_and_team(slide):
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            if "@SIH" in shape.text:
                shape.text_frame.clear()
            if "Your Team Name" in shape.text:
                format_text(shape, "EIDOLON", 12, True, PP_ALIGN.CENTER)

# Slide 1: TITLE PAGE
s1 = prs.slides[0]
for shape in s1.shapes:
    if not hasattr(shape, "text"): continue
    t = shape.text
    if "TITLE PAGE" in t:
        format_text(shape, "EIDOLON // SENTINEL-NET", 28, True)
    elif "SMART INDIA HACKATHON" in t:
        format_text(shape, "SMART INDIA HACKATHON 2026", 24, True)
    elif "Problem Statement ID" in t:
        txt = ("Problem Statement ID: SIH26145\n"
               "Problem Statement Title: AI-Based Detection of Cyber Threats in Unidirectional IP Traffic\n"
               "Theme: Software / Blockchain & Cybersecurity\n"
               "PS Category: Software\n"
               "Organization: NTRO\n"
               "Team Name: EIDOLON")
        format_text(shape, txt, 16)

# Slide 2: PROPOSED SOLUTION
s2 = prs.slides[1]
clean_footer_and_team(s2)
for shape in s2.shapes:
    if not hasattr(shape, "text"): continue
    t = shape.text
    if "IDEA TITLE" in t or "EIDOLON" in t:
        format_text(shape, "EIDOLON // SENTINEL-NET", 32, True)
    elif "Proposed Solution" in t or "Problem:" in t:
        txt = ("THE PROBLEM:\n"
               "• Unidirectional IP traffic only: no return path, no active probes, no mitigation commands.\n"
               "• Detection must rely entirely on passive packet/flow metadata.\n\n"
               "PROPOSED SOLUTION:\n"
               "• EIDOLON: A passive, read-only, AI-driven, explainable cyber-threat detection platform.\n"
               "• Flow aggregation combined with 52 canonical behavioral features.\n"
               "• 'Observe. Analyze. Explain.'")
        format_text(shape, txt, 16)

# Slide 3: TECHNICAL APPROACH
s3 = prs.slides[2]
clean_footer_and_team(s3)
for shape in s3.shapes:
    if not hasattr(shape, "text"): continue
    t = shape.text
    if "TECHNICAL APPROACH" in t:
        format_text(shape, "TECHNICAL APPROACH", 32, True)
    elif "Technologies" in t or "Architecture Pipeline" in t:
        # We will clear this text box to make room for the PowerPoint shape diagram
        shape.text_frame.clear()

# Draw architecture diagram using PPT shapes on Slide 3
arch_steps = [
    "Traffic/PCAP", "Packet Parsing", "Flow Aggregation", "52 Features", 
    "Preprocessing", "ML Detection", "Evidence/Risk", "SOC Command Center"
]
start_y = 2.0
y_step = 0.55
for i, step in enumerate(arch_steps):
    # Box
    box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(start_y + i*y_step), Inches(4), Inches(0.4))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0, 112, 192) # SIH blueish
    box.line.color.rgb = RGBColor(255, 255, 255)
    format_text(box, step, 14, True, PP_ALIGN.CENTER, (255,255,255))
    # Down arrow (except last)
    if i < len(arch_steps) - 1:
        arrow = s3.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(2.4), Inches(start_y + i*y_step + 0.4), Inches(0.2), Inches(0.15))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(100, 100, 100)
        arrow.line.fill.background()

# AI details on the right of Slide 3
ai_box = s3.shapes.add_textbox(Inches(5.0), Inches(2.0), Inches(4.5), Inches(3.0))
format_text(ai_box, "AI & FEATURE ENGINEERING:\n\n• 52 canonical features (41 usable after preprocessing).\n• Models: XGBoost, Random Forest, Logistic Regression, Isolation Forest.\n• Simulated unidirectional telemetry constraint (passive/read-only).\n• Supervised models detect learned classes; Isolation Forest flags anomalies.", 16)
ai_box.text_frame.word_wrap = True

# Slide 4: FEASIBILITY AND VIABILITY
s4 = prs.slides[3]
clean_footer_and_team(s4)
for shape in s4.shapes:
    if not hasattr(shape, "text"): continue
    t = shape.text
    if "FEASIBILITY AND VIABILITY" in t:
        format_text(shape, "FEASIBILITY AND VIABILITY", 32, True)
    elif "Analysis" in t or "SCIENTIFIC" in t:
        txt = ("SCIENTIFIC HONESTY & LIMITATIONS:\n"
               "• Unseen C2 Recall: 0% (@ 1% FPR) and 1.7% (@ 5% FPR).\n"
               "• Supervised models fail when classes are entirely unseen. We do not claim universal zero-day detection.\n"
               "• Anomaly detection provides novelty-ranking, but operational FPR remains high.\n\n"
               "ROBUSTNESS & EXPLAINABILITY:\n"
               "• System resists timing jitter, packet-size noise, duration scaling, and metadata dropout.\n"
               "• ATT&CK mapping serves as contextual evidence, not definitive proof of attribution.\n"
               "• System does NOT claim to be production-ready.")
        format_text(shape, txt, 16)
        shape.text_frame.word_wrap = True

# Slide 5: IMPACT AND BENEFITS
s5 = prs.slides[4]
clean_footer_and_team(s5)
for shape in s5.shapes:
    if not hasattr(shape, "text"): continue
    t = shape.text
    if "IMPACT AND BENEFITS" in t:
        format_text(shape, "IMPACT AND BENEFITS", 32, True)
    elif "Potential" in t or "VERIFIED" in t:
        txt = ("VERIFIED EVALUATION METRICS (Phase 7/8):\n"
               "• CICIDS2017 dataset: 2,313,810 flows, 8 scenarios.\n"
               "• Supervised Baseline: XGBoost Macro-F1 = 0.2675; Logistic Regression Macro-F1 = 0.3389\n"
               "• Strongest Generalization (Unseen DDoS): ROC-AUC = 0.887, PR-AUC = 0.776, Anomaly F1 = 0.571\n"
               "• DoS-Wednesday LOSO F1: 0.806\n"
               "• Cross-scenario Macro-F1: 0.197 ± 0.033\n\n"
               "IMPACT:\n"
               "• Protects highly segmented/air-gapped networks using passive read-only capability.\n"
               "• Gives SOC analysts a transparent decision surface without risking reverse-path exposure.")
        format_text(shape, txt, 16)
        shape.text_frame.word_wrap = True

# Slide 6: RESEARCH AND REFERENCES
s6 = prs.slides[5]
clean_footer_and_team(s6)
for shape in s6.shapes:
    if not hasattr(shape, "text"): continue
    t = shape.text
    if "RESEARCH" in t:
        format_text(shape, "DEMO & RESEARCH", 32, True)
    elif "Details" in t or "SCIENTIFIC" in t:
        txt = ("DETERMINISTIC PCAP REPLAY / RECORDED TRAFFIC:\n"
               "• The demo utilizes recorded traffic through the real ingestion pipeline, NOT live Internet traffic.\n"
               "• End-to-End: PCAP → Pipeline → Feature Extraction → ML Inference → SQLite → WebSocket → SOC.\n\n"
               "SCIENTIFIC RIGOR:\n"
               "• 531 automated tests pass locally.\n"
               "• Real-data evaluation with explicit known vs. unseen splits.")
        format_text(shape, txt, 16)
        shape.text_frame.word_wrap = True

# Embed REAL Screenshot on Slide 6 (Demo & Research)
try:
    s6.shapes.add_picture('real_soc.png', Inches(4.5), Inches(3.5), width=Inches(5))
except Exception as e:
    print(f"Screenshot insertion failed: {e}")

# Remove instruction slide if it exists
xml_slides = prs.slides._sldIdLst  
slides = list(xml_slides)
if len(slides) > 6:
    xml_slides.remove(slides[6])

prs.save('EIDOLON_SIH2026_FINAL.pptx')
print("Successfully generated EIDOLON_SIH2026_FINAL.pptx")
