from pptx import Presentation
prs = Presentation('SIH2026-IDEA-Presentation-Format.pptx')
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for j, shape in enumerate(slide.shapes):
        if hasattr(shape, "text_frame") and shape.text_frame:
            print(f"Shape {j} ({shape.name}): {shape.text[:100].replace(chr(10), ' ')}")
