from pptx import Presentation
from pptx.util import Pt

prs = Presentation('SIH2026-IDEA-Presentation-Format.pptx')

# We know the slides:
# Slide 0: Title
# Slide 1: Proposed Solution
# Slide 2: Technical Approach
# Slide 3: Feasibility and Viability
# Slide 4: Impact and Benefits
# Slide 5: Research and References
# Slide 6: Instructions (will delete)

# Slide 0 (Title)
slide0 = prs.slides[0]
for shape in slide0.shapes:
    if not shape.has_text_frame: continue
    text = shape.text.lower()
    if 'problem statement id' in text:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = "Problem Statement ID: SIH26145\nProblem Statement Title: AI-Based Detection of Cyber Threats in Unidirectional IP Traffic\nTheme: Software / Blockchain & Cybersecurity\nPS Category: Software\nOrganization: NTRO"
        p.font.size = Pt(16)

# Slide 1 (Proposed Solution)
slide1 = prs.slides[1]
for shape in slide1.shapes:
    if not shape.has_text_frame: continue
    text = shape.text.lower()
    if 'idea title' in text:
        shape.text_frame.text = "EIDOLON // SENTINEL-NET"
    elif 'proposed solution' in text:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = "Problem:\n- AI-based detection of cyber threats in unidirectional IP traffic.\n- Constraint: Passive/read-only observation. No probes, no return-path dependency.\n\nOur Solution: EIDOLON\n- Passive, read-only, AI-driven, explainable, near-real-time architecture.\n- Replaces active polling with deterministic flow aggregation & scenario-aware ML inference on unidirectional metadata."
        p.font.size = Pt(16)

# Slide 2 (Technical Approach)
slide2 = prs.slides[2]
for shape in slide2.shapes:
    if not shape.has_text_frame: continue
    text = shape.text.lower()
    if 'technical approach' in text:
        pass # Keep title
    elif 'technologies' in text:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = "Architecture Pipeline:\nTraffic / PCAP → Packet Parsing → Flow Aggregation → 52 Canonical Features → Preprocessing → ML Models → SOC Command Center via REST/WebSocket\n\nAI & Feature Engineering:\n- Models: XGBoost, Random Forest, Logistic Regression, Isolation Forest\n- 52 canonical features (41 usable post-preprocessing). Highlights: flow behavior, packet size, IAT, TCP flags, directionality.\n- Explainability: MITRE ATT&CK contextual mapping & deterministic rationale."
        p.font.size = Pt(16)

# Slide 3 (Feasibility and Viability)
slide3 = prs.slides[3]
for shape in slide3.shapes:
    if not shape.has_text_frame: continue
    text = shape.text.lower()
    if 'feasibility' in text:
        pass
    elif 'analysis' in text:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = "Feasibility & Production Readiness:\n- Testing: 531 tests passing (483 backend, 48 frontend) proving system stability.\n- Robustness: Controlled perturbation experiments showed low classification change under timing jitter, packet noise, duration variation, and metadata dropout.\n\nCritical Scientific Limitations Analyzed:\n- C2 Detection Challenge: Unseen C2 received 0-1.7% recall under strict scenario-aware test split (operational FPR 1-5%).\n- Note: Supervised + anomaly fusion did not improve this specific weakness."
        p.font.size = Pt(16)

# Slide 4 (Impact and Benefits)
slide4 = prs.slides[4]
for shape in slide4.shapes:
    if not shape.has_text_frame: continue
    text = shape.text.lower()
    if 'impact and benefits' in text:
        pass
    elif 'potential' in text:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = "Performance & Generalization Results:\n- Supervised: XGBoost macro-F1 (0.2675); Logistic Regression macro-F1 (0.3389).\n- Anomaly & Novel Threat Findings: Isolation Forest provides crucial novelty-ranking signal.\n- Unseen DDoS: ROC-AUC 0.887, PR-AUC 0.776, Anomaly F1 0.571 (strongest generalization case).\n- DoS-Wednesday LOSO F1: 0.806.\n\nTarget Use-Cases & Impact:\n- Critical infrastructure monitoring and air-gapped/segmented environments.\n- Empowers SOC analyst workflows through a seamless UI."
        p.font.size = Pt(16)

# Slide 5 (Research and References)
slide5 = prs.slides[5]
for shape in slide5.shapes:
    if not shape.has_text_frame: continue
    text = shape.text.lower()
    if 'research' in text:
        pass
    elif 'details' in text:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = "Scientific Evaluation Methodology:\n- Dataset: CICIDS2017 (2,313,810 flows, 8 scenarios).\n- Process: Scenario-aware evaluation to prevent data leakage (random IID splitting avoided). Explicit known vs unseen evaluation.\n\nDemo & SOC Command Center Workflow:\n- DETERMINISTIC PCAP REPLAY — NOT LIVE CAPTURE.\n- Recorded PCAP → ML Inference → DetectionEvent → WebSocket → SOC Dashboard.\n- Real-time display of Overview, Detections, Flows, Sensor status, and Threat context."
        p.font.size = Pt(16)

# Delete slide 6 (Instructions)
xml_slides = prs.slides._sldIdLst  
slides = list(xml_slides)
xml_slides.remove(slides[6])

prs.save('SIH2026_Final_Presentation.pptx')
print("Successfully created SIH2026_Final_Presentation.pptx")
