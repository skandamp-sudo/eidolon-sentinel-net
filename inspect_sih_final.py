from pptx import Presentation
prs = Presentation('SIH2026_Final_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}")
bg_color = prs.slides[0].background.fill.fore_color if prs.slides[0].background.fill.type else None
print(f"Slide 0 background: {bg_color}")
