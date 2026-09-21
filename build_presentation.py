import os
import matplotlib.pyplot as plt
import networkx as nx
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- DARK THEME SETTINGS ---
BG_COLOR = RGBColor(18, 20, 26)
TEXT_COLOR = RGBColor(230, 230, 240)
ACCENT_BLUE = RGBColor(0, 190, 255)
ACCENT_GREEN = RGBColor(0, 255, 140)
ACCENT_RED = RGBColor(255, 80, 80)
MUTED_TEXT = RGBColor(150, 150, 160)

# --- HELPER FUNCTIONS ---
def set_dark_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_title(slide, text):
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.33), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    return title_box

def add_subtitle(slide, text, top=1.2):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.33), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.color.rgb = MUTED_TEXT
    return box

def add_bullet_points(slide, points, left=0.5, top=2.0, width=12.33, height=5.0, font_size=22):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, pt_text in enumerate(points):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = pt_text
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT_COLOR
        p.level = 0
        p.space_after = Pt(14)
    return box

def create_bar_chart():
    plt.style.use('dark_background')
    fig, ax = plt.subplots(figsize=(7, 4))
    fig.patch.set_facecolor('#12141a')
    ax.set_facecolor('#12141a')
    
    metrics = ['XGBoost\nMacro-F1', 'LogReg\nMacro-F1', 'DDoS\nROC-AUC', 'DDoS\nPR-AUC', 'DoS-Wed\nLOSO F1']
    values = [0.2675, 0.3389, 0.887, 0.776, 0.806]
    colors = ['#00beff', '#00beff', '#00ff8c', '#00ff8c', '#00ff8c']
    
    bars = ax.bar(metrics, values, color=colors, edgecolor='none')
    ax.set_ylim(0, 1.0)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#555566')
    ax.spines['bottom'].set_color('#555566')
    
    for bar in bars:
        yval = bar.get_height()
        ax.text(bar.get_x() + bar.get_width()/2, yval + 0.02, f'{yval:.4f}', ha='center', va='bottom', color='white', fontweight='bold')
    
    plt.tight_layout()
    plt.savefig('chart_results.png', dpi=200, facecolor=fig.get_facecolor(), transparent=True)
    plt.close()

def create_arch_diagram():
    plt.style.use('dark_background')
    fig, ax = plt.subplots(figsize=(10, 3.5))
    fig.patch.set_facecolor('#12141a')
    ax.set_facecolor('#12141a')
    
    nodes = ["PCAP\nTraffic", "Passive\nIngestion", "Flow\nAggregation", "52 Features\nExtraction", "AI Inference\n(Supervised+IF)", "Evidence\n+ Risk", "SOC\nDashboard"]
    G = nx.DiGraph()
    for i in range(len(nodes)-1):
        G.add_edge(nodes[i], nodes[i+1])
        
    pos = {n: (i, 0) for i, n in enumerate(nodes)}
    nx.draw_networkx_edges(G, pos, edge_color='#00beff', arrows=True, arrowsize=20, width=2)
    
    # Draw custom nodes
    for k, v in pos.items():
        box_color = '#1f2430'
        ax.text(v[0], v[1], k, ha='center', va='center', color='white',
                bbox=dict(facecolor=box_color, edgecolor='#00beff', boxstyle='round,pad=0.5', linewidth=1.5),
                fontsize=10, fontweight='bold')
        
    ax.set_xlim(-0.5, len(nodes)-0.5)
    ax.set_ylim(-0.5, 0.5)
    ax.axis('off')
    plt.tight_layout()
    plt.savefig('chart_arch.png', dpi=200, facecolor=fig.get_facecolor(), transparent=True)
    plt.close()

def create_ui_mockup():
    img = Image.new('RGB', (1600, 900), color=(18, 20, 26))
    d = ImageDraw.Draw(img)
    # Header
    d.rectangle([(0,0), (1600, 60)], fill=(30, 35, 45))
    d.text((20, 20), "EIDOLON // SOC COMMAND CENTER", fill=(0, 190, 255))
    # Main panels
    d.rectangle([(20, 80), (1100, 500)], fill=(25, 30, 38), outline=(0, 190, 255), width=2)
    d.text((40, 100), "LIVE EVENT STREAM", fill=(255, 255, 255))
    d.rectangle([(1120, 80), (1580, 500)], fill=(25, 30, 38))
    d.text((1140, 100), "THREAT DISTRIBUTION", fill=(255, 255, 255))
    # Bottom panels
    d.rectangle([(20, 520), (530, 880)], fill=(25, 30, 38))
    d.text((40, 540), "ACTIVE FLOWS", fill=(255, 255, 255))
    d.rectangle([(550, 520), (1580, 880)], fill=(25, 30, 38), outline=(0, 255, 140), width=1)
    d.text((570, 540), "DETECTION EVIDENCE & ATT&CK CONTEXT", fill=(0, 255, 140))
    # Draw some mock bars
    for i in range(5):
        d.rectangle([(1140, 150 + i*60), (1140 + (150-i*20), 180 + i*60)], fill=(0, 190, 255))
    img.save('mock_soc.png')

# Generate visual assets
create_bar_chart()
create_arch_diagram()
create_ui_mockup()

# --- INITIALIZE PPTX ---
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# --- SLIDE 1: TITLE ---
s1 = prs.slides.add_slide(blank_layout)
set_dark_bg(s1)
title = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
p = title.text_frame.paragraphs[0]
p.text = "EIDOLON // SENTINEL-NET"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER

sub = s1.shapes.add_textbox(Inches(1), Inches(3.8), Inches(11.33), Inches(1))
p2 = sub.text_frame.paragraphs[0]
p2.text = "Passive AI Threat Detection for Unidirectional IP Traffic"
p2.font.size = Pt(28)
p2.font.color.rgb = TEXT_COLOR
p2.alignment = PP_ALIGN.CENTER

info = s1.shapes.add_textbox(Inches(1), Inches(6.0), Inches(11.33), Inches(1))
p3 = info.text_frame.paragraphs[0]
p3.text = "Team: EIDOLON\nSIH26145 | NTRO | Software & Cybersecurity"
p3.font.size = Pt(20)
p3.font.color.rgb = MUTED_TEXT
p3.alignment = PP_ALIGN.CENTER

# --- SLIDE 2: THE PROBLEM ---
s2 = prs.slides.add_slide(blank_layout)
set_dark_bg(s2)
add_title(s2, "THE PROBLEM")
add_subtitle(s2, "Core Constraint: Unidirectional Observation")
add_bullet_points(s2, [
    "Traffic can be observed only in one direction.",
    "No return path. No active probes. No mitigation commands.",
    "Detection must rely entirely on passive packet/flow metadata."
], top=2.2)
# Visual
box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(11.33), Inches(1.5))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(30, 35, 45)
box.line.color.rgb = ACCENT_RED
tf = box.text_frame
tf.text = "NETWORK  ⟶  ONE-WAY TELEMETRY  ⟶  SENSOR"
tf.paragraphs[0].font.size = Pt(32)
tf.paragraphs[0].font.color.rgb = TEXT_COLOR
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# --- SLIDE 3: OUR SOLUTION ---
s3 = prs.slides.add_slide(blank_layout)
set_dark_bg(s3)
add_title(s3, "OUR SOLUTION")
add_subtitle(s3, "Observe. Analyze. Explain.")
add_bullet_points(s3, [
    "Passive Ingestion: Real-time, strictly passive listening.",
    "Flow Aggregation: Reassembling metadata without state injection.",
    "Behavioral Analysis: Extracting 52 canonical flow features.",
    "AI Detection: Supervised classification + anomaly ranking.",
    "Explainable Evidence: Transparent decision surfaces for SOC analysts."
], top=2.2)

# --- SLIDE 4: SYSTEM ARCHITECTURE ---
s4 = prs.slides.add_slide(blank_layout)
set_dark_bg(s4)
add_title(s4, "SYSTEM ARCHITECTURE")
add_subtitle(s4, "End-to-End Passive Pipeline")
s4.shapes.add_picture('chart_arch.png', Inches(0.5), Inches(2.5), width=Inches(12.33))

# --- SLIDE 5: AI + FEATURE ENGINEERING ---
s5 = prs.slides.add_slide(blank_layout)
set_dark_bg(s5)
add_title(s5, "AI + FEATURE ENGINEERING")
add_bullet_points(s5, [
    "52 Canonical Behavioral Features (41 usable after preprocessing)",
    "Feature Groups: Basic, Directional, Packet Size, IAT, TCP Flags, Protocol, Payload",
    "Supervised Models: XGBoost, Random Forest, Logistic Regression",
    "Unsupervised: Isolation Forest for behavioral novelty-ranking",
    "Strategy: Supervised models identify learned classes; Anomaly detection ranks unknown behaviors."
], top=2.0)

# --- SLIDE 6: SOC COMMAND CENTER ---
s6 = prs.slides.add_slide(blank_layout)
set_dark_bg(s6)
add_title(s6, "SOC COMMAND CENTER")
add_subtitle(s6, "Live Event Stream • Active Flows • Threat Distribution • Evidence")
s6.shapes.add_picture('mock_soc.png', Inches(1.66), Inches(2.0), width=Inches(10))

# --- SLIDE 7: SCIENTIFIC EVALUATION ---
s7 = prs.slides.add_slide(blank_layout)
set_dark_bg(s7)
add_title(s7, "SCIENTIFIC EVALUATION")
add_subtitle(s7, "Scenario-Aware Methodology on Real Data")
add_bullet_points(s7, [
    "Dataset: CICIDS2017",
    "Scale: 2,313,810 flows across 8 distinct attack scenarios",
    "Features: 52 canonical features",
    "Reliability: 531 automated verification tests",
    "Methodology: Strict scenario-aware evaluation to test true generalization, rejecting flawed random IID splitting."
], top=2.2)

# --- SLIDE 8: KEY RESULTS ---
s8 = prs.slides.add_slide(blank_layout)
set_dark_bg(s8)
add_title(s8, "KEY RESULTS")
add_subtitle(s8, "Phase 7 & 8 Evaluation")
s8.shapes.add_picture('chart_results.png', Inches(5.5), Inches(2.0), width=Inches(7))
add_bullet_points(s8, [
    "Unseen DDoS:",
    "  • ROC-AUC = 0.887",
    "  • PR-AUC = 0.776",
    "  • F1 = 0.571",
    "",
    "Phase 7 Baseline:",
    "  • XGBoost Macro-F1 = 0.2675",
    "  • LogReg Macro-F1 = 0.3389",
    "",
    "Cross-Scenario Macro-F1:",
    "  • 0.197 ± 0.033"
], top=2.0, width=5.0, font_size=20)

# --- SLIDE 9: NOVEL THREAT FINDINGS ---
s9 = prs.slides.add_slide(blank_layout)
set_dark_bg(s9)
add_title(s9, "NOVEL THREAT FINDINGS")
add_subtitle(s9, "Scientific Rigor: Evaluating the C2 Limitation")
add_bullet_points(s9, [
    "Unseen C2 Detection Challenge:",
    "  • Recall @ 1% FPR = 0%",
    "  • Recall @ 5% FPR = 1.7%",
    "",
    "Key Finding: Supervised models fail when a class is entirely unseen.",
    "Anomaly detectors provide a ranking signal, but high false positive rates complicate operational classification.",
    "",
    "Scientific Integrity: We do not claim universal zero-day detection."
], top=2.2)

# --- SLIDE 10: ROBUSTNESS + EXPLAINABILITY ---
s10 = prs.slides.add_slide(blank_layout)
set_dark_bg(s10)
add_title(s10, "ROBUSTNESS & EXPLAINABILITY")
add_bullet_points(s10, [
    "Controlled Perturbation Experiments:",
    "  • Timing jitter, packet-size noise, duration scaling",
    "  • Directional imbalance, metadata dropout",
    "  • Result: Demonstrated low classification volatility.",
    "",
    "Explainable Pipeline:",
    "  • Detection ⟶ Evidence ⟶ Risk/Severity ⟶ ATT&CK Context",
    "  • ATT&CK mapping is contextual evidence, not proof of attribution."
], top=2.2)

# --- SLIDE 11: LIVE DEMO / WORKFLOW ---
s11 = prs.slides.add_slide(blank_layout)
set_dark_bg(s11)
add_title(s11, "DEMO WORKFLOW")
add_subtitle(s11, "Deterministic PCAP Replay / Recorded Traffic")
box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.5), Inches(11.33), Inches(4))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(25, 30, 38)
box.line.color.rgb = ACCENT_BLUE
tf = box.text_frame
tf.word_wrap = True
tf.text = "1. Deterministic PCAP Replay (Recorded Traffic, Not Live Internet)\n2. Real ingestion pipeline & flow aggregation\n3. Feature extraction\n4. Real ML inference\n5. Database persistence\n6. WebSocket streaming\n7. SOC Command Center dashboard rendering"
for p in tf.paragraphs:
    p.font.size = Pt(24)
    p.font.color.rgb = TEXT_COLOR
    p.space_after = Pt(14)
    p.alignment = PP_ALIGN.LEFT

# --- SLIDE 12: IMPACT + FUTURE ---
s12 = prs.slides.add_slide(blank_layout)
set_dark_bg(s12)
add_title(s12, "IMPACT & FUTURE")
add_bullet_points(s12, [
    "Target Applications:",
    "  • One-way monitoring environments & Air-gapped networks",
    "  • Critical infrastructure security operations",
    "  • Passive network telemetry analysis",
    "",
    "Future Work:",
    "  • Better threshold calibration & novel-threat algorithms",
    "  • Broader datasets with diverse real-world traffic",
    "  • Real unidirectional hardware sensor validation",
    "  • Operational deployment hardening"
], top=2.0)

# --- SLIDE 13: CLOSING ---
s13 = prs.slides.add_slide(blank_layout)
set_dark_bg(s13)
t = s13.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11.33), Inches(1.5))
p = t.text_frame.paragraphs[0]
p.text = "EIDOLON DOESN'T FIGHT THE NETWORK.\nIT OBSERVES IT."
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.alignment = PP_ALIGN.CENTER

sub = s13.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11.33), Inches(2))
p2 = sub.text_frame.paragraphs[0]
p2.text = "Observe what is available.\nExtract behavioral signals.\nDetect known and anomalous activity.\nExplain the evidence.\nGive the analyst a decision surface."
p2.font.size = Pt(24)
p2.font.color.rgb = TEXT_COLOR
p2.alignment = PP_ALIGN.CENTER

footer = s13.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.33), Inches(0.5))
p3 = footer.text_frame.paragraphs[0]
p3.text = "Team: EIDOLON  |  SIH26145  |  NTRO"
p3.font.size = Pt(18)
p3.font.color.rgb = MUTED_TEXT
p3.alignment = PP_ALIGN.CENTER

prs.save("EIDOLON_SIH2026_Final.pptx")
print("Presentation generated successfully: EIDOLON_SIH2026_Final.pptx")
