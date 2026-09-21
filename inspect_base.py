from pptx import Presentation

try:
    prs = Presentation('SIH2026_Final_Presentation.pptx')
    print(f"Total slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                print(f"TEXT: {shape.text[:100].replace(chr(10), ' ')}")
except Exception as e:
    print("Error:", e)
