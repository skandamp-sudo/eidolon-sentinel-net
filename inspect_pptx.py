from pptx import Presentation

prs = Presentation('SIH2026-IDEA-Presentation-Format.pptx')
print(f"Total slides in template: {len(prs.slides)}")

print("\nAvailable layouts:")
for i, layout in enumerate(prs.slide_layouts):
    print(f"Layout {i}: {layout.name}")

print("\nSlide contents:")
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(f"TEXT: {shape.text[:100].replace(chr(10), ' ')}")
