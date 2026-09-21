from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation('SIH2026-IDEA-Presentation-Format.pptx')

# Dark theme colors
bg_color = RGBColor(18, 20, 26)
text_color = RGBColor(230, 230, 240)
accent_blue = RGBColor(0, 190, 255)

def apply_dark_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

def format_text(shape, text, font_size=14, color=text_color, bold=False):
    if shape.has_text_frame:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        if "Title" in shape.name:
            p.font.color.rgb = accent_blue

# Slide 1
s0 = prs.slides[0]
apply_dark_bg(s0)
for shape in s0.shapes:
    if "Subtitle" in shape.name:
        format_text(shape, "EIDOLON // SENTINEL-NET", 32, accent_blue, True)
    elif "Title" in shape.name:
        format_text(shape, "SMART INDIA HACKATHON 2026", 24, text_color, True)
    elif "TextBox" in shape.name:
        txt = ("Problem Statement ID: SIH26145\n"
               "Problem Statement Title: AI-Based Detection of Cyber Threats in Unidirectional IP Traffic\n"
               "Theme: Software / Blockchain & Cybersecurity\n"
               "PS Category: Software\n"
               "Organization: NTRO\n"
               "Team: EIDOLON")
        format_text(shape, txt, 18, text_color)

# Function to clean footers
def clean_footer_and_team(slide):
    for shape in slide.shapes:
        if "Footer" in shape.name or (hasattr(shape, "text") and "@SIH" in shape.text):
            shape.text_frame.clear()
        if "Oval" in shape.name or (hasattr(shape, "text") and "Your Team Name" in shape.text):
            format_text(shape, "EIDOLON", 12, bg_color, True)
            if hasattr(shape, "fill"):
                shape.fill.solid()
                shape.fill.fore_color.rgb = accent_blue

# Slide 2 (Solution)
s1 = prs.slides[1]
apply_dark_bg(s1)
clean_footer_and_team(s1)
for shape in s1.shapes:
    if "Title" in shape.name:
        format_text(shape, "EIDOLON // SENTINEL-NET", 32, accent_blue, True)
    elif "TextBox" in shape.name:
        txt = ("THE PROBLEM:\n"
               "• Traffic can be observed only in one direction. No return path. No active probes.\n"
               "• Detection must rely entirely on passive packet/flow metadata.\n\n"
               "OUR SOLUTION (EIDOLON):\n"
               "• Passive, read-only, AI-driven, explainable, near-real-time architecture.\n"
               "• Replaces active polling with deterministic flow aggregation & scenario-aware ML inference.\n"
               "• 'Observe. Analyze. Explain.'")
        format_text(shape, txt, 16, text_color)
s1.shapes.add_picture('visual_problem.png', Inches(1), Inches(4.5), width=Inches(8))

# Slide 3 (Approach)
s2 = prs.slides[2]
apply_dark_bg(s2)
clean_footer_and_team(s2)
for shape in s2.shapes:
    if "Title" in shape.name:
        format_text(shape, "TECHNICAL APPROACH", 32, accent_blue, True)
    elif "TextBox" in shape.name:
        txt = ("SYSTEM ARCHITECTURE & PIPELINE:\n"
               "• Deterministic PCAP Replay → Passive Ingestion → Flow Aggregation → 52 Canonical Features\n"
               "• Preprocessing → Supervised ML + Anomaly Detection → Evidence/Risk → WebSocket → SOC Dashboard\n\n"
               "AI & FEATURE ENGINEERING:\n"
               "• 52 Canonical behavioral features (41 usable). Groups: Directional, IAT, TCP Flags, Payload, Size.\n"
               "• Models: XGBoost, Random Forest, Logistic Regression, Isolation Forest.")
        format_text(shape, txt, 15, text_color)
s2.shapes.add_picture('visual_pipeline.png', Inches(0.5), Inches(4.2), width=Inches(9))

# Slide 4 (Feasibility)
s3 = prs.slides[3]
apply_dark_bg(s3)
clean_footer_and_team(s3)
for shape in s3.shapes:
    if "Title" in shape.name:
        format_text(shape, "FEASIBILITY AND VIABILITY", 32, accent_blue, True)
    elif "TextBox" in shape.name:
        txt = ("SCIENTIFIC HONESTY & LIMITATIONS:\n"
               "• Unseen C2 Detection Challenge: Recall @ 1% FPR = 0%; Recall @ 5% FPR = 1.7%.\n"
               "• Supervised models fail on unseen classes. We do not claim universal zero-day detection.\n"
               "• Anomaly detectors provide novelty-ranking, but operational FPR remains a challenge.\n\n"
               "ROBUSTNESS & EXPLAINABILITY:\n"
               "• Low classification change under timing jitter, packet noise, duration scaling, & metadata dropout.\n"
               "• ATT&CK mapping provides contextual evidence, not definitive proof of attribution.")
        format_text(shape, txt, 15, text_color)
s3.shapes.add_picture('visual_novelty.png', Inches(1), Inches(4.5), width=Inches(8))

# Slide 5 (Impact)
s4 = prs.slides[4]
apply_dark_bg(s4)
clean_footer_and_team(s4)
for shape in s4.shapes:
    if "Title" in shape.name:
        format_text(shape, "IMPACT AND BENEFITS", 32, accent_blue, True)
    elif "TextBox" in shape.name:
        txt = ("VERIFIED METRICS (Phase 7 & 8):\n"
               "• Supervised Baseline: XGBoost Macro-F1 = 0.2675, LogReg Macro-F1 = 0.3389\n"
               "• Strongest Generalization: Unseen DDoS (ROC-AUC = 0.887, PR-AUC = 0.776, Anomaly F1 = 0.571)\n"
               "• DoS-Wednesday LOSO F1 = 0.806. Cross-scenario Macro-F1 = 0.197 ± 0.033\n\n"
               "IMPACT & FUTURE POTENTIAL:\n"
               "• Critical infrastructure monitoring and air-gapped/segmented environments.\n"
               "• Empowers SOC analyst workflows through a seamless UI.\n"
               "• EIDOLON doesn't fight the network. It observes it.")
        format_text(shape, txt, 14, text_color)
        shape.width = Inches(5.5) # Make room for chart
s4.shapes.add_picture('visual_results.png', Inches(5.8), Inches(2.2), width=Inches(4.0))

# Slide 6 (Research)
s5 = prs.slides[5]
apply_dark_bg(s5)
clean_footer_and_team(s5)
for shape in s5.shapes:
    if "Title" in shape.name:
        format_text(shape, "RESEARCH AND REFERENCES", 32, accent_blue, True)
    elif "TextBox" in shape.name:
        txt = ("SCIENTIFIC EVALUATION METHODOLOGY:\n"
               "• Dataset: CICIDS2017 (2,313,810 flows, 8 scenarios, 52 canonical features).\n"
               "• Scale: 531 automated verification tests (Backend & Frontend).\n"
               "• Methodology: Scenario-aware evaluation to prevent data leakage (random IID splitting rejected).\n\n"
               "DEMO & SOC COMMAND CENTER WORKFLOW:\n"
               "• DETERMINISTIC PCAP REPLAY — NOT LIVE CAPTURE.\n"
               "• Recorded PCAP → ML Inference → DetectionEvent → WebSocket → SOC Dashboard.\n"
               "• Live event stream, threat distribution, active flows, and contextual mapping.")
        format_text(shape, txt, 16, text_color)

# Remove slide 7
xml_slides = prs.slides._sldIdLst  
slides = list(xml_slides)
if len(slides) > 6:
    xml_slides.remove(slides[6])

prs.save('EIDOLON_SIH2026_FINAL.pptx')
print("Successfully created EIDOLON_SIH2026_FINAL.pptx")
