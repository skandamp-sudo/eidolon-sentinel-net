import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
from pptx import Presentation
from pptx.util import Inches

# 1. Generate Architecture Diagram
plt.figure(figsize=(8, 4))
G = nx.DiGraph()
nodes = ["Traffic\n(PCAP)", "Packet\nParsing", "Flow\nAggregation", "52 Features\n(Preprocessing)", "ML Inference\n(XGBoost+IForest)", "SOC Command\nCenter"]
for i in range(len(nodes)-1):
    G.add_edge(nodes[i], nodes[i+1])

pos = {n: (i, 0) for i, n in enumerate(nodes)}
nx.draw(G, pos, with_labels=True, node_color='lightblue', edge_color='gray', node_size=3000, font_size=8, font_weight='bold', arrows=True, arrowsize=20)
plt.title("EIDOLON // SENTINEL-NET Passive Architecture", pad=20)
plt.tight_layout()
plt.savefig("arch_diagram.png", dpi=150)
plt.close()

# 2. Generate Results Chart
plt.figure(figsize=(6, 4))
labels = ['XGBoost\nMacro-F1', 'LogReg\nMacro-F1', 'DDoS\nROC-AUC', 'DDoS\nPR-AUC', 'DoS-Wed\nLOSO F1']
values = [0.2675, 0.3389, 0.887, 0.776, 0.806]
colors = ['#1f77b4', '#1f77b4', '#ff7f0e', '#ff7f0e', '#2ca02c']

bars = plt.bar(labels, values, color=colors)
plt.ylim(0, 1.0)
plt.title("Phase 7 & 8 Evaluation Highlights")
plt.ylabel("Score")
for bar in bars:
    yval = bar.get_height()
    plt.text(bar.get_x() + bar.get_width()/2.0, yval + 0.02, round(yval, 4), ha='center', va='bottom', fontweight='bold')
plt.tight_layout()
plt.savefig("results_chart.png", dpi=150)
plt.close()

# Insert into PPTX
prs = Presentation('SIH2026_Final_Presentation.pptx')

# Slide 2 (Technical Approach) -> Insert Architecture Diagram
slide2 = prs.slides[2]
# Find a place to put the image (e.g., bottom half)
slide2.shapes.add_picture("arch_diagram.png", Inches(1), Inches(4), width=Inches(8))

# Slide 4 (Impact and Benefits) -> Insert Results Chart
slide4 = prs.slides[4]
slide4.shapes.add_picture("results_chart.png", Inches(4.5), Inches(3), width=Inches(5))

prs.save('SIH2026_Final_Presentation.pptx')
print("Visuals added successfully.")
