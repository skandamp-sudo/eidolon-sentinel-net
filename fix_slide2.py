from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation('EIDOLON_SIH2026_FINAL.pptx')
slide2 = prs.slides[1]

def format_text(shape, text, font_size=None, bold=None, align=None, color_rgb=None):
    if shape.has_text_frame:
        shape.text_frame.clear()
        p = shape.text_frame.paragraphs[0]
        p.text = text
        if font_size: p.font.size = Pt(font_size)
        if bold is not None: p.font.bold = bold
        if align: p.alignment = align
        if color_rgb: p.font.color.rgb = RGBColor(*color_rgb)

# Identify title and content shapes
for shape in slide2.shapes:
    if not shape.has_text_frame:
        continue
    # Usually Title is at the top
    if shape.top < Inches(1.5):
        format_text(shape, "THE PROBLEM", 32, True)
    # The main content box is usually larger and lower
    elif shape.top > Inches(1.5) and shape.width > Inches(4):
        # We'll use this for the text content
        txt = ("Modern threat detection often assumes bidirectional network visibility.\n\n"
               "But SIH26145 imposes a fundamentally different constraint:\n\n"
               "• Traffic can be observed only in one direction.\n"
               "• No return-path dependency.\n"
               "• No active probes or scanning.\n"
               "• No mitigation commands can be sent back.\n"
               "• Detection must rely entirely on passively observable packet/flow metadata.")
        format_text(shape, txt, 16)
        # Adjust width to make room for diagram on the right
        shape.width = Inches(7)
        shape.height = Inches(4.5)
        shape.text_frame.word_wrap = True

# Highlight Statement at the bottom
highlight_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(6.0), Inches(7), Inches(0.8))
highlight_box.fill.solid()
highlight_box.fill.fore_color.rgb = RGBColor(240, 240, 240) # Light grey highlight
highlight_box.line.color.rgb = RGBColor(0, 112, 192) # SIH blue border
format_text(highlight_box, "“HOW DO WE DETECT CYBER THREATS WHEN WE CAN ONLY OBSERVE?”", 16, True, PP_ALIGN.CENTER, (0, 112, 192))

# --- Diagram on the right ---
start_x = Inches(8.5)
start_y = Inches(1.8)
y_step = Inches(1.2)
box_w = Inches(3.5)
box_h = Inches(0.6)

blocks = [
    "OBSERVABLE NETWORK",
    "ONE-WAY IP TRAFFIC",
    "EIDOLON SENSOR",
    "AI THREAT DETECTION"
]

for i, text in enumerate(blocks):
    box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, start_x, start_y + i*y_step, box_w, box_h)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0, 112, 192) # SIH blueish
    box.line.color.rgb = RGBColor(255, 255, 255)
    format_text(box, text, 14, True, PP_ALIGN.CENTER, (255,255,255))
    
    # Down arrow
    if i < len(blocks) - 1:
        arrow = slide2.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, start_x + Inches(1.6), start_y + i*y_step + box_h + Inches(0.05), Inches(0.3), Inches(0.45))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(100, 100, 100)
        arrow.line.color.rgb = RGBColor(100, 100, 100)
        
# Show reverse direction as unavailable
# Draw an up arrow next to the first down arrow (Network <- One-Way Traffic)
up_arrow = slide2.shapes.add_shape(MSO_SHAPE.UP_ARROW, start_x + Inches(0.8), start_y + box_h + Inches(0.05), Inches(0.3), Inches(0.45))
up_arrow.fill.solid()
up_arrow.fill.fore_color.rgb = RGBColor(220, 220, 220) # Faded
up_arrow.line.color.rgb = RGBColor(200, 200, 200)
# Red cross over it
cross = slide2.shapes.add_shape(MSO_SHAPE.MATH_MULTIPLY, start_x + Inches(0.75), start_y + box_h + Inches(0.075), Inches(0.4), Inches(0.4))
cross.fill.solid()
cross.fill.fore_color.rgb = RGBColor(255, 0, 0)
cross.line.color.rgb = RGBColor(255, 0, 0)

# Add small text "NO RETURN PATH" next to the crossed arrow
no_return = slide2.shapes.add_textbox(start_x - Inches(0.3), start_y + box_h, Inches(1.5), Inches(0.4))
format_text(no_return, "No Return Path", 10, True, PP_ALIGN.CENTER, (255, 0, 0))

prs.save('EIDOLON_SIH2026_FINAL.pptx')
print("Slide 2 successfully fixed.")
